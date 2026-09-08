from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Lingowave_Proje_Durum_Raporu.docx"

BLUE = "1F4E79"
LIGHT_BLUE = "EAF2F8"
PALE_BLUE = "F5F9FC"
GRID = "D9E2F3"
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = GRID, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 110, start: int = 120, bottom: int = 110, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_widths(table, widths: list[float]) -> None:
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)


def style_run(run, *, bold: bool = False, size: float = 10.3, color: RGBColor = BLACK, italic: bool = False) -> None:
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def style_paragraph(paragraph, *, space_after: float = 6, line: float = 1.08) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = line


def add_text(doc, text: str, *, bold_lead: str | None = None, size: float = 10.3, after: float = 6) -> None:
    p = doc.add_paragraph()
    style_paragraph(p, space_after=after)
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        style_run(r, bold=True, size=size)
        r2 = p.add_run(text[len(bold_lead) :])
        style_run(r2, size=size)
    else:
        r = p.add_run(text)
        style_run(r, size=size)


def add_bullet(doc, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    style_paragraph(p, space_after=3, line=1.04)
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    r = p.add_run(text)
    style_run(r, size=9.8)


def add_heading(doc, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    style_run(r, bold=True, size=15.0 if level == 1 else 11.5)


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[float], font_size: float = 8.8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_widths(table, widths)
    header = table.rows[0]
    repeat_header(header)
    prevent_row_split(header)
    for cell, text in zip(header.cells, headers):
        set_cell_shading(cell, BLUE)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style_paragraph(p, space_after=0, line=1.0)
        r = p.add_run(text)
        style_run(r, bold=True, size=font_size, color=RGBColor(255, 255, 255))
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        if row_index % 2 == 1:
            fill = PALE_BLUE
        else:
            fill = "FFFFFF"
        for cell, text in zip(row.cells, values):
            set_cell_shading(cell, fill)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            style_paragraph(p, space_after=0, line=1.0)
            r = p.add_run(text)
            style_run(r, size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_key_value_table(doc, rows: list[tuple[str, str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = widths or [2.35, 4.85]
    set_widths(table, widths)
    header = table.rows[0]
    repeat_header(header)
    prevent_row_split(header)
    for cell, label in zip(header.cells, ("Alan", "Değer")):
        set_cell_shading(cell, BLUE)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        style_paragraph(p, space_after=0, line=1.0)
        r = p.add_run(label)
        style_run(r, bold=True, size=8.9, color=RGBColor(255, 255, 255))
    for i, (key, value) in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for cell in row.cells:
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, PALE_BLUE if i % 2 else "FFFFFF")
        p0, p1 = row.cells[0].paragraphs[0], row.cells[1].paragraphs[0]
        style_paragraph(p0, space_after=0, line=1.0)
        style_paragraph(p1, space_after=0, line=1.0)
        r0 = p0.add_run(key)
        style_run(r0, bold=True, size=8.9)
        r1 = p1.add_run(value)
        style_run(r1, size=8.9)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10.3)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.color.rgb = BLACK


def configure_section(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.70)
    section.right_margin = Inches(0.70)
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.55)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)


def add_footer(section) -> None:
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, space_after=0, line=1.0)
    r = p.add_run("LingoWave | Proje durum ve üretim doğrulama | 2026-09-08")
    style_run(r, size=8.5, color=RGBColor(90, 90, 90))


def build() -> None:
    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        configure_section(section)
        add_footer(section)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(7)
    title.paragraph_format.keep_with_next = True
    r = title.add_run("LingoWave Proje Durum ve Üretim Doğrulama Raporu")
    style_run(r, bold=True, size=22)

    meta = doc.add_paragraph()
    style_paragraph(meta, space_after=10, line=1.0)
    rm = meta.add_run("Durum tarihi: 8 Eylül 2026  |  Dal: codex/production-saas  |  Commit: 18898e9")
    style_run(rm, size=10.1, color=RGBColor(75, 75, 75))

    add_text(
        doc,
        "Sonuç: LingoWave’in serverless üretim yolu, gerçek CPU medya işleme doğrulaması ve yeni frontend workbench dağıtımı tamamlandı. Canlı akış CloudFront → API Gateway → Lambda → Aurora Data API → S3/SQS → CPU worker → S3 output şeklinde çalışıyor; gerçek dublaj çıktısı indirildi ve doğrulandı. CPU worker test sonrasında 0/0/0’a döndü, GPU kapasitesi sıfırda tutuldu ve açık GPU kota talebi performans yükseltmesi olarak korunuyor.",
        size=10.7,
        after=9,
    )

    add_heading(doc, "Yönetici Özeti", 1)
    add_table(
        doc,
        ["Konu", "Durum", "Kanıt / not"],
        [
            ["Serverless üretim API", "TAMAM", "CloudFront → API Gateway → Lambda → Aurora Data API; canlı /health HTTP 200"],
            ["Gerçek CPU dubbing E2E", "TAMAM", "API → S3 → SQS → CPU worker → S3 → indirilebilir MP4"],
            ["Google/deep-translator", "VARSAYILAN", "Normal çeviri yolu; AWS Translate yalnızca opsiyonel"],
            ["Hy-MT2 refinement", "SEÇİCİ", "Süre eşiği aşılırsa en fazla bir pass/segment"],
            ["Yeni frontend UI/UX", "CANLI", "Commit 18898e9; CloudFront asset ve invalidation doğrulandı"],
            ["Google OAuth", "TAMAM", "Gerçek login, callback, yeni kullanıcı, linking, session, logout ve cancel"],
            ["GPU kapasitesi", "BEKLEMEDE", "Kota talebi açık; worker/ASG 0/0/0 ve CPU yolu bloklanmıyor"],
        ],
        [2.15, 1.05, 4.0],
    )

    add_heading(doc, "Son Rapor Güncellemesinden Sonra Yapılanlar", 1)
    for item in [
        "Frontend translation workbench yenilendi: Upload → Configure → Export akışı, source media paneli, audio overview, Translation/Voice/Subtitles/Advanced inspector sekmeleri, cost/start paneli ve recent projects görünümü production’a aktarıldı.",
        "Commit 18898e9, origin/codex/production-saas dalına push edildi. Terraform yalnızca frontend S3 hashed asset’lerini ve index.html’i güncelledi; API, worker, GPU, queue, database, IAM ve billing kaynaklarına değişiklik yapılmadı.",
        "CloudFront distribution E29DUJBS3MCQ75 için I7B4X1VJUQPYQ0QJFW326AJO4D invalidation tamamlandı. Canlı HTML yeni index-BvQHbPr6.js ve index-DI96JvC1.css asset’lerini sunuyor.",
        "Chrome canlı doğrulamasında gerçek authenticated workspace, kullanıcı kredileri, Voice/Advanced/Translation sekmeleri, recent projects ve responsive davranış kontrol edildi; browser error/warning log’u görülmedi.",
        "Rapor ve operasyon belgeleri 8 Eylül 2026 canlı kanıtlarıyla hizalandı. Repository çalışma ağacı temiz ve remote ile senkron.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "Canlı Mimari ve Üretim Pipeline", 1)
    add_text(doc, "Production’da doğrulanan tam akış:", size=10.2, after=4)
    add_table(
        doc,
        ["Aşama", "Provider / davranış", "Durum"],
        [
            ["1", "CloudFront → API Gateway → Lambda kontrol düzlemi", "CANLI / sağlıklı"],
            ["2", "Presigned S3 upload ve S3/SQS job akışı", "CANLI / doğrulandı"],
            ["3", "Aurora PostgreSQL Serverless v2 Data API", "CANLI / auto-pause 0 ACU"],
            ["4", "CPU Fargate Spot worker; queue-driven 0 → 1 → 0", "CANLI / ölçek-sıfır"],
            ["5", "Demucs → Whisper → GoogleTranslator", "Gerçek CPU E2E PASS"],
            ["6", "Opsiyonel Hy-MT2; süre eşiği ve tek pass", "Routing benchmark PASS"],
            ["7", "VoxCPM2 TTS → bounded FFmpeg timing/mix", "Gerçek CPU E2E PASS"],
            ["8", "Private S3 output → job status → signed download", "Gerçek çıktı ve ffprobe PASS"],
        ],
        [0.55, 4.8, 1.85],
    )
    add_text(doc, "Üretim voice provider’ı openbmb/VoxCPM2’dir. Google/deep-translator varsayılan çeviri sağlayıcısıdır; AWS Translate normal dublaj yolunda zorunlu değildir. Hy-MT2 yalnızca süre toleransı aşıldığında seçici refinement olarak devreye girer ve segment başına en fazla bir kez çalışır.", size=10.1)

    add_heading(doc, "Frontend Production Durumu", 1)
    add_key_value_table(
        doc,
        [
            ("Public URL", "https://d3ncg3eqih0ccj.cloudfront.net"),
            ("CloudFront distribution", "E29DUJBS3MCQ75"),
            ("Frontend deploy commit", "18898e9 — feat: refresh translation workbench UI"),
            ("Invalidation", "I7B4X1VJUQPYQ0QJFW326AJO4D — Completed"),
            ("Live asset check", "HTTP 200; index-BvQHbPr6.js ve index-DI96JvC1.css"),
            ("Live health", "HTTP 200; database configured, storage s3, queue sqs"),
            ("Browser verification", "Authenticated workbench, tabs, recent projects, no console errors/warnings"),
            ("Responsive verification", "390×844; horizontal overflow yok"),
        ],
    )

    add_heading(doc, "Gerçek CPU E2E Ölçümleri", 1)
    add_text(doc, "8 Eylül 2026’da üretim CloudFront API’si üzerinden gönderilen 13.2 saniyelik gerçek medya varlığı ile tam yol doğrulandı. Queue alarmı CPU Fargate Spot worker’ı 0’dan 1’e çıkardı; prewarmed image çekildi, worker işi aldı, gerçek Demucs, Whisper, GoogleTranslator, VoxCPM2 ve FFmpeg çalıştı, MP4 private S3’e yüklendi ve imzalı çıktı indirildi. İş tamamlanınca kapasite 0/0/0’a döndü.", size=10.1)
    add_table(
        doc,
        ["Metrik", "Ölçüm", "Açıklama"],
        [
            ["Input duration", "13.2 s", "Gerçek medya"],
            ["Total processing", "93.4195 s", "Worker işlem süresi"],
            ["Queue wait", "330.3478 s", "Image pull/startup dahil"],
            ["Demucs", "7.9456 s", "Speech/background separation"],
            ["Whisper", "13.3249 s", "Timestamped transcription"],
            ["Translation", "0.1246 s", "1 segment Google translation"],
            ["Hy-MT2", "0 s", "Bu E2E’de tetiklenmedi"],
            ["VoxCPM2", "66.2701 s stage wall", "53.9630 s synthesis telemetry"],
            ["Generated audio", "9.76 s", "48 kHz WAV validation"],
            ["FFmpeg / mix", "0.4425 s", "Timing, mix, mux"],
            ["Upload", "0.2709 s", "S3 output upload"],
            ["Peak RAM / CPU", "11,272.195 MB / 131.979%", "Worker telemetry"],
            ["RTF", "5.5290 / 7.0772", "VoxCPM2 / whole job"],
            ["Estimated cost", "$0.006046", "Input dakika başına $0.027482"],
        ],
        [1.55, 1.85, 3.8],
        font_size=8.6,
    )

    add_heading(doc, "Çeviri ve Süre Politikası", 1)
    add_table(
        doc,
        ["Metrik", "Sonuç"],
        [
            ["Toplam / Google / Hy-MT2 segment", "1 / 1 / 0"],
            ["Refinement rate", "0%"],
            ["Average deviation before / after", "−26.0606% / −26.0606%"],
            ["Google translation time", "0.1246 s"],
            ["Separate real Hy-MT2 CPU benchmark", "156.5701 s / 1 segment"],
            ["Targeted routing benchmark", "3 case; fit refinement yok, 2 mismatch tek pass"],
            ["Final speed adjustment", "Bounded FFmpeg; text truncate edilmedi"],
        ],
        [2.7, 4.5],
        font_size=9.0,
    )
    add_text(doc, "İşlem sırası doğal çeviri → ilk TTS → duration comparison → gerekirse tek Hy-MT2 refinement → ikinci TTS → son bounded speed adjustment şeklindedir. Her segment için original duration, first/refined TTS duration, deviation, refinement kullanımı, refined text ve final speed ratio telemetry’ye yazılır.", size=10.1)

    add_heading(doc, "AWS, Maliyet ve Güvenlik Durumu", 1)
    add_key_value_table(
        doc,
        [
            ("CPU worker", "desired/running/pending 0/0/0; queue-driven target live"),
            ("GPU worker / ASG", "0/0/0; architecture and pinned image retained"),
            ("GPU quota", "CASE_OPENED; request cancelled değil; quota 0"),
            ("Expensive compute", "Şu an çalışmıyor"),
            ("Aurora", "0 ACU auto-pause; Data API canlı"),
            ("$25 budget guardrail", "Değiştirilmedi"),
            ("ECR", "API 4 manifest / 1.04 GiB; worker 4 manifest / 15.90 GiB logical"),
            ("OAuth secrets", "Değerler rapora, log’a veya commit’e yazılmadı"),
        ],
    )
    add_text(doc, "İdle compute maliyeti sıfıra yakın tutuluyor. ECR, Aurora storage/backups, snapshot, Secrets Manager, CloudWatch, S3 ve CloudFront gibi kullanım/depoya bağlı kalemler devam edebilir. Tam aylık Cost Explorer faturası normal billing gecikmesi sonrasında ayrıca doğrulanmalıdır.", size=10.1)

    add_heading(doc, "Testler ve Doğrulamalar", 1)
    add_table(
        doc,
        ["Kontrol", "Durum", "Not"],
        [
            ["Backend tests", "PASS", "Final live migration gate ve provider/E2E kanıtı"],
            ["Frontend production build", "PASS", "Vite build başarılı"],
            ["Ruff", "PASS", "Full backend kapsamı"],
            ["Bandit", "PASS", "CI medium-severity threshold; düşük bulgular bilgilendirici"],
            ["pip-audit / npm audit", "PASS", "Bilinen zafiyet bulunmadı"],
            ["Terraform fmt / validate / plan / apply", "PASS", "Frontend apply sonrası plan değişiklik göstermedi"],
            ["Live /health", "PASS", "CloudFront HTTP 200"],
            ["Live browser QA", "PASS", "UI, tabs, auth workspace, responsive, no console errors"],
            ["Local Playwright", "KISMİ", "5 GPU case skip; 4 local signup testi localhost:8000 backend yokluğu nedeniyle tamamlanmadı"],
            ["Repository", "CLEAN / PUSHED", "origin/codex/production-saas ile senkron"],
        ],
        [2.05, 1.0, 4.15],
        font_size=8.7,
    )

    add_heading(doc, "Açık Konular ve Kullanıcı Kararları", 1)
    for item in [
        "GPU kota talebi açık bırakıldı; CPU production yolu bundan bağımsız çalışıyor. Kota onaylandığında VoxCPM2 throughput/latency yükseltmesi için GPU worker yolu kullanılabilir.",
        "CPU worker concurrency varsayılanı maliyet ve hesap vCPU kotası nedeniyle 1’de tutuluyor. Daha yüksek paralellik ancak canlı yük ölçümü ve quota artışı sonrası yükseltilmeli.",
        "Model ve lisans incelemesi yapılmadı. Kullanıcı, production öncesi VoxCPM2, Demucs, Whisper/WhisperX, Hy-MT2 ve diğer provider lisanslarını ayrıca inceleyecek.",
        "Tam aylık idle maliyet denetimi, AWS billing aggregation tamamlandıktan sonra yapılmalı.",
        "Stripe ödeme akışı deployment secret’ları ve price ID’ler sağlanana kadar bilerek kapalı tutuluyor; uygulama başarılı ödeme simüle etmiyor.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "Final Status Fields", 1)
    add_key_value_table(
        doc,
        [
            ("DEFAULT TRANSLATION PROVIDER", "Google Translate via deep-translator"),
            ("REFINEMENT PROVIDER", "tencent/Hy-MT2-1.8B; selective, maximum one pass"),
            ("TTS / VOICE PROVIDER", "openbmb/VoxCPM2; pinned revision; 48 kHz output"),
            ("REAL FULL DUBBING E2E", "PASS"),
            ("FINAL DUBBED MEDIA GENERATED", "YES"),
            ("OUTPUT DOWNLOAD VERIFIED", "YES"),
            ("DATABASE MIGRATION", "0013_google_oauth_identities applied"),
            ("GOOGLE AUTH", "PASS — new user, linked login, safe linking, password login, session, logout, cancel"),
            ("SECRETS EXPOSED", "NO"),
            ("LIVE API", "PASS — CloudFront → API Gateway → Lambda; no always-on API ECS"),
            ("FRONTEND PRODUCTION", "PASS — deployed and live-verified"),
            ("LICENSE REVIEW", "NOT PERFORMED — USER WILL REVIEW BEFORE PRODUCTION"),
            ("MANUAL USER ACTION STILL REQUIRED", "None for current deployment; license and billing review remain before commercial launch"),
        ],
    )

    add_text(doc, "Kaynak raporlar: docs/AWS_MEDIA_PIPELINE_STATUS.md, docs/serverless-migration-gates.md, docs/aws-cost-audit.md ve docs/ecr-retention-audit.md. Bu Word raporu, bu belgelerdeki 8 Eylül 2026 canlı kanıtlarını tek bir yönetici raporunda birleştirir.", size=8.9, after=0)

    props = doc.core_properties
    props.title = "LingoWave Proje Durum ve Üretim Doğrulama Raporu"
    props.subject = "Production serverless, CPU E2E, frontend deployment and live verification"
    props.author = "LingoWave Engineering"
    props.last_modified_by = "LingoWave Engineering"
    props.comments = "Updated 2026-09-08 from live deployment and repository evidence."
    props.modified = datetime.now(timezone.utc)
    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
